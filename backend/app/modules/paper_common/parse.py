"""论文应用 · 文档解析与分块。

支持：
- PDF  (PyMuPDF / fitz)         —— 已装在运行环境
- DOCX (python-docx)            —— 已装在运行环境
- URL  (httpx + 标准库 html.parser) —— 零额外依赖抽取正文

并提供 clean_text / chunk_text 用于大文件分块处理与检索召回。
"""
from __future__ import annotations

import html
import httpx
import re
import uuid
from html.parser import HTMLParser
from pathlib import Path
from typing import Optional

# backend/uploads/<code>/ 存放上传原文，便于追溯与下载
UPLOAD_ROOT = Path(__file__).resolve().parents[3] / "uploads"


def ensure_upload_dir(code: str) -> Path:
    d = UPLOAD_ROOT / code
    d.mkdir(parents=True, exist_ok=True)
    return d


def save_upload(code: str, filename: str, content: bytes) -> tuple[str, Path]:
    """保存上传文件，返回 (file_id, 落盘路径)。"""
    d = ensure_upload_dir(code)
    ext = Path(filename).suffix.lower()
    fid = uuid.uuid4().hex
    path = d / f"{fid}{ext}"
    path.write_bytes(content)
    return fid, path


def extract_pdf(path: Path) -> str:
    import fitz

    doc = fitz.open(str(path))
    try:
        pages = [page.get_text("text") for page in doc]
    finally:
        doc.close()
    return "\n\n".join(p.strip() for p in pages if p.strip())


def extract_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    out: list[str] = []
    for p in doc.paragraphs:
        if p.text and p.text.strip():
            out.append(p.text.strip())
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                out.append(" | ".join(cells))
    return "\n".join(out)


class _HTMLTextExtractor(HTMLParser):
    """极简正文抽取：跳过 script/style/head，收集段落级标签文本。"""

    SKIP = {"script", "style", "head", "meta", "link", "noscript"}
    COLLECT = {"p", "article", "li", "h1", "h2", "h3", "h4", "h5", "h6", "blockquote", "td", "section"}

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skip_depth = 0
        self._buf: list[str] = []

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self._skip_depth += 1
        if tag in self.COLLECT:
            self._flush()

    def handle_endtag(self, tag):
        if tag in self.SKIP and self._skip_depth > 0:
            self._skip_depth -= 1
        if tag in self.COLLECT:
            self._flush()

    def handle_data(self, data):
        if self._skip_depth == 0:
            self._buf.append(data)

    def _flush(self):
        text = "".join(self._buf).strip()
        if text:
            self.parts.append(text)
        self._buf = []


def extract_url(url: str, timeout: float = 25.0) -> str:
    resp = httpx.get(
        url,
        timeout=timeout,
        follow_redirects=True,
        headers={"User-Agent": "Mozilla/5.0 (compatible; BaizePaper/1.0)"},
    )
    resp.raise_for_status()
    parser = _HTMLTextExtractor()
    parser.feed(resp.text)
    text = "\n\n".join(parser.parts)
    # 兜底：若正文过少，直接去标签
    if len(text) < 200:
        text = re.sub(r"<[^>]+>", " ", resp.text)
        text = html.unescape(text)
        text = re.sub(r"\s+", " ", text).strip()
    return text


def clean_text(text: str) -> str:
    text = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def chunk_text(text: str, size: int = 1400) -> list[str]:
    """按段落切分为约 size 字符的块（单段超长则按句切），用于大文件分块处理与检索。"""
    text = clean_text(text)
    if not text:
        return []
    if len(text) <= size:
        return [text]
    paras = [p for p in re.split(r"\n+", text) if p.strip()]
    chunks: list[str] = []
    cur = ""
    for p in paras:
        if len(cur) + len(p) + 1 <= size:
            cur = (cur + "\n" + p).strip()
        else:
            if cur:
                chunks.append(cur)
            if len(p) > size:
                chunks.extend(_split_long(p, size))
                cur = ""
            else:
                cur = p
    if cur:
        chunks.append(cur)
    return chunks


def _split_long(p: str, size: int) -> list[str]:
    parts = re.split(r"(?<=[。.!?！？;；])", p)
    out: list[str] = []
    cur = ""
    for s in parts:
        if len(cur) + len(s) <= size:
            cur += s
        else:
            if cur:
                out.append(cur)
            cur = s
    if cur:
        out.append(cur)
    return out
